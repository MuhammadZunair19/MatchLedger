import csv
import io
import os
import re
from decimal import Decimal
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape
from fastapi import HTTPException
from sqlalchemy import select
from .. import models as m
from ..config import settings
from . import review_service


def safe_text(value):
    text = '' if value is None else str(value)
    return "'" + text if text.startswith(('=', '+', '-', '@')) else text


def report_snapshot(db, org, match_run_id):
    run = db.scalar(select(m.MatchRun).where(m.MatchRun.organization_id == org, m.MatchRun.id == match_run_id))
    if not run: raise HTTPException(404, 'Matching run not found')
    invoice = db.scalar(select(m.SupplierInvoice).where(m.SupplierInvoice.organization_id == org, m.SupplierInvoice.id == run.invoice_id))
    order = db.scalar(select(m.PurchaseOrder).where(m.PurchaseOrder.organization_id == org, m.PurchaseOrder.id == run.purchase_order_id))
    supplier = db.scalar(select(m.Supplier).where(m.Supplier.organization_id == org, m.Supplier.id == invoice.supplier_id))
    receipts = db.scalars(select(m.GoodsReceipt).where(m.GoodsReceipt.organization_id == org, m.GoodsReceipt.id.in_(run.receipt_ids or []))).all()
    findings = db.scalars(select(m.MatchFinding).where(m.MatchFinding.organization_id == org,
        m.MatchFinding.match_run_id == run.id).order_by(m.MatchFinding.created_at)).all()
    approval = db.scalar(select(m.InvoiceApproval).where(m.InvoiceApproval.organization_id == org,
        m.InvoiceApproval.invoice_id == invoice.id))
    decisions = db.scalars(select(m.ApprovalDecision).where(m.ApprovalDecision.organization_id == org,
        m.ApprovalDecision.invoice_id == invoice.id).order_by(m.ApprovalDecision.created_at)).all()
    approval_value = review_service.public(approval) if approval else {'status': 'DRAFT'}
    if approval:
        submitted = db.get(m.User, approval.submitted_by) if approval.submitted_by else None
        decided = db.get(m.User, approval.decided_by) if approval.decided_by else None
        approval_value['submitted_name'] = submitted.name if submitted else None
        approval_value['decided_name'] = decided.name if decided else None
    decision_values = []
    for decision in decisions:
        value = review_service.public(decision); reviewer = db.get(m.User, decision.reviewer_id)
        value['reviewer_name'] = reviewer.name if reviewer else None; decision_values.append(value)
    finding_rows = []
    for item in findings:
        review = db.scalar(select(m.ExceptionReview).where(m.ExceptionReview.organization_id == org,
            m.ExceptionReview.match_finding_id == item.id))
        value = review_service.public(item)
        value['resolution_status'] = review.status if review else item.resolution_status
        value['resolution_reason'] = review.resolution_reason if review else None
        finding_rows.append(value)
    return {'report_version': '1.0', 'match_run': review_service.public(run), 'summary': run.result_summary or {},
        'invoice': review_service.public(invoice), 'purchase_order': review_service.public(order),
        'supplier': review_service.public(supplier), 'receipts': [review_service.public(row) for row in receipts],
        'findings': finding_rows, 'approval': approval_value,
        'decisions': decision_values, 'source_versions': run.source_record_versions,
        'document_versions': run.document_versions}


def _rtl(value):
    text = safe_text(value)
    if not re.search(r'[\u0600-\u06ff]', text): return escape(text)
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return escape(get_display(arabic_reshaper.reshape(text)))
    except ImportError:
        return escape(text)


def build_pdf(snapshot, organization_name, generated_at):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    font_paths = [r'C:\Windows\Fonts\arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    bold_paths = [r'C:\Windows\Fonts\arialbd.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    regular = next((path for path in font_paths if os.path.exists(path)), None)
    bold = next((path for path in bold_paths if os.path.exists(path)), regular)
    font_name = 'Helvetica'; bold_name = 'Helvetica-Bold'
    if regular:
        pdfmetrics.registerFont(TTFont('MatchLedgerUnicode', regular)); font_name = 'MatchLedgerUnicode'
        pdfmetrics.registerFont(TTFont('MatchLedgerUnicodeBold', bold)); bold_name = 'MatchLedgerUnicodeBold'
    buffer = io.BytesIO(); styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MLTitle', parent=styles['Title'], fontName=bold_name, textColor=colors.HexColor('#123044'), fontSize=20))
    styles.add(ParagraphStyle(name='MLBody', parent=styles['BodyText'], fontName=font_name, fontSize=8.5, leading=11))
    styles.add(ParagraphStyle(name='MLRight', parent=styles['MLBody'], alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name='MLSmall', parent=styles['MLBody'], fontSize=7, textColor=colors.HexColor('#64748b')))
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=14*mm, leftMargin=14*mm, topMargin=14*mm, bottomMargin=14*mm)
    story = [Paragraph('MatchLedger', styles['MLTitle']), Paragraph('Matching and approval report', styles['Heading2']),
        Paragraph(f"Report ID: {snapshot['match_run']['id']} | Generated: {generated_at} | Organization: {_rtl(organization_name)}", styles['MLSmall']), Spacer(1, 5*mm)]
    summary=snapshot['summary']; invoice=snapshot['invoice']; order=snapshot['purchase_order']; supplier=snapshot['supplier']
    details = [['Invoice', _rtl(invoice['number']), 'Invoice date', invoice['date'], 'Due date', invoice.get('due_date') or 'Not recorded'],
        ['Purchase order', order['number'], 'Supplier', _rtl(supplier['name']), 'Currency', invoice['currency']],
        ['Receipts', ', '.join(row['number'] for row in snapshot['receipts']) or 'None', 'Mode', snapshot['match_run']['matching_mode'], 'Match status', snapshot['match_run']['status']],
        ['Invoice value', summary.get('invoice_value'), 'Ordered value', summary.get('agreed_order_value'), 'Supported value', summary.get('supported_value')],
        ['Price variance', summary.get('price_variance'), 'Quantity variance', summary.get('quantity_variance'), 'Difference', summary.get('difference_against_supported')]]
    table = Table([[Paragraph(_rtl(cell), styles['MLBody']) for cell in row] for row in details], colWidths=[24*mm, 54*mm, 27*mm, 54*mm, 25*mm, 54*mm])
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.HexColor('#cbd5e1')),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#e8f2ee')),
        ('BACKGROUND',(2,0),(2,-1),colors.HexColor('#e8f2ee')),('BACKGROUND',(4,0),(4,-1),colors.HexColor('#e8f2ee')),('VALIGN',(0,0),(-1,-1),'TOP'),('FONTNAME',(0,0),(-1,-1),font_name),('FONTSIZE',(0,0),(-1,-1),8)]))
    story += [table, Spacer(1,5*mm), Paragraph('Line item comparison', styles['Heading2'])]
    headers=['Product / SKU','PO qty','Received','Invoiced','PO unit price','Invoice unit price','Supported qty','Difference']
    rows=[headers]
    for item in summary.get('comparisons', []):
        rows.append([_rtl(f"{item.get('product','')} / {item.get('sku') or '-'}"), item.get('po_quantity'), item.get('received_quantity'), item.get('invoice_quantity'),
            item.get('po_unit_price'), item.get('invoice_unit_price'), item.get('supported_quantity'), item.get('line_difference')])
    line_table=Table([[Paragraph(str(cell or ''), styles['MLBody']) for cell in row] for row in rows], repeatRows=1,
        colWidths=[65*mm,22*mm,22*mm,22*mm,28*mm,30*mm,25*mm,25*mm])
    line_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#123044')),('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),.25,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fafc')])]))
    story += [line_table, PageBreak(), Paragraph('Exceptions and resolution', styles['Heading2'])]
    exc_rows=[['Finding','Severity','Expected','Actual','Difference','Resolution','Explanation / reason']]
    for item in snapshot['findings']:
        exc_rows.append([item['code'],item['severity'],item.get('expected_value'),item.get('actual_value'),item.get('difference'),
            item.get('resolution_status'),_rtl(f"{item.get('explanation') or ''} {item.get('resolution_reason') or ''}".strip())])
    exc_table=Table([[Paragraph(str(cell or ''), styles['MLBody']) for cell in row] for row in exc_rows], repeatRows=1,
        colWidths=[35*mm,22*mm,25*mm,25*mm,25*mm,32*mm,85*mm])
    exc_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#7f1d1d')),('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),.25,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP')]))
    approval=snapshot['approval']
    story += [exc_table, Spacer(1,5*mm), Paragraph('Approval', styles['Heading2']),
        Paragraph(f"Status: {approval.get('status','DRAFT')} | Reviewer: {_rtl(approval.get('decided_name') or approval.get('submitted_name') or 'Not assigned')} | Decision date: {approval.get('decided_at') or 'Not decided'} | Reason: {_rtl(approval.get('decision_reason') or 'None')}", styles['MLBody']),
        Spacer(1,5*mm), Paragraph(f"Report version 1.0 | Generated by MatchLedger | Source snapshot: {escape(safe_text(snapshot.get('source_versions')))}", styles['MLSmall'])]
    def footer(canvas, _doc):
        canvas.saveState(); canvas.setFont(font_name,7); canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawString(14*mm,7*mm,'Generated by MatchLedger'); canvas.drawRightString(283*mm,7*mm,f'Page {canvas.getPageNumber()}'); canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def build_matching_excel(snapshot):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); ws=wb.active; ws.title='Summary'
    summary=snapshot['summary']; invoice=snapshot['invoice']; order=snapshot['purchase_order']
    summary_rows=[('Invoice number',safe_text(invoice['number'])),('Invoice date',invoice['date']),('Due date',invoice.get('due_date')),
        ('Purchase order',safe_text(order['number'])),('Supplier',safe_text(snapshot['supplier']['name'])),('Currency',invoice['currency']),
        ('Match status',snapshot['match_run']['status']),('Invoice value',Decimal(summary.get('invoice_value','0'))),
        ('Ordered value',Decimal(summary.get('agreed_order_value','0'))),('Supported value',Decimal(summary.get('supported_value','0'))),
        ('Difference',Decimal(summary.get('difference_against_supported','0'))),('Approval status',snapshot['approval'].get('status','DRAFT'))]
    for row in summary_rows: ws.append(row)
    ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=45; ws.freeze_panes='A2'
    for cell in ws['A']: cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='E8F2EE')
    for row in range(8,12): ws.cell(row,2).number_format='#,##0.00'
    lines=wb.create_sheet('Line Item Comparison'); headers=['Product','SKU','PO Quantity','Received Quantity','Invoice Quantity','PO Unit Price','Invoice Unit Price','Supported Quantity','Difference','Currency']
    lines.append(headers)
    for item in summary.get('comparisons',[]): lines.append([safe_text(item.get('product')),safe_text(item.get('sku')),Decimal(item.get('po_quantity','0')),Decimal(item.get('received_quantity','0')),Decimal(item.get('invoice_quantity','0')),Decimal(item.get('po_unit_price','0')),Decimal(item.get('invoice_unit_price','0')),Decimal(item.get('supported_quantity','0')),Decimal(item.get('line_difference','0')),safe_text(summary.get('currency'))])
    exceptions=wb.create_sheet('Exceptions'); exceptions.append(['Code','Severity','Expected','Actual','Difference','Unit','Resolution','Resolution Reason','Explanation'])
    for item in snapshot['findings']: exceptions.append([safe_text(item['code']),safe_text(item['severity']),safe_text(item.get('expected_value')),safe_text(item.get('actual_value')),safe_text(item.get('difference')),safe_text(item.get('unit')),safe_text(item.get('resolution_status')),safe_text(item.get('resolution_reason')),safe_text(item.get('explanation'))])
    approvals=wb.create_sheet('Approval History'); approvals.append(['Decision','Reviewer ID','Reason','Timestamp','Invoice Version'])
    for item in snapshot['decisions']: approvals.append([safe_text(item['decision']),safe_text(item.get('reviewer_name') or item['reviewer_id']),safe_text(item.get('reason')),item['created_at'],safe_text(item['invoice_version'])])
    refs=wb.create_sheet('Audit References'); refs.append(['Match Run ID','Input Fingerprint','Document Versions','Source Record Versions']); refs.append([snapshot['match_run']['id'],snapshot['match_run'].get('input_fingerprint'),safe_text(snapshot['document_versions']),safe_text(snapshot['source_versions'])])
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='123044')
        for column in sheet.columns:
            letter=column[0].column_letter; sheet.column_dimensions[letter].width=min(60,max(12,max(len(str(cell.value or '')) for cell in column)+2))
    buffer=io.BytesIO(); wb.save(buffer); return buffer.getvalue()


def build_register_excel(headers, rows, title='Register'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); ws=wb.active; ws.title=title[:31]; ws.append(headers)
    for row in rows: ws.append([safe_text(value) if isinstance(value,str) else value for value in row])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='123044')
    for column in ws.columns: ws.column_dimensions[column[0].column_letter].width=min(60,max(12,max(len(str(cell.value or '')) for cell in column)+2))
    buffer=io.BytesIO(); wb.save(buffer); return buffer.getvalue()


def build_csv(headers, rows):
    buffer=io.StringIO(); writer=csv.writer(buffer); writer.writerow(headers)
    for row in rows: writer.writerow([safe_text(value) for value in row])
    return ('\ufeff'+buffer.getvalue()).encode('utf-8')


def persist_export(db, ctx, report_type, related_id, filename, content_type, content, snapshot):
    root=Path(settings.report_storage_path).resolve(); root.mkdir(parents=True,exist_ok=True)
    row=m.ReportExport(organization_id=ctx.org,created_by=ctx.user.id,updated_by=ctx.user.id,report_type=report_type,
        related_entity_id=related_id,status='PROCESSING',filename=filename,content_type=content_type,snapshot=snapshot)
    db.add(row); db.flush(); path=(root/f'{ctx.org}_{row.id}_{uuid4().hex}').resolve()
    if root not in path.parents: raise HTTPException(500,'Invalid report storage path')
    path.write_bytes(content); row.storage_key=str(path); row.status='COMPLETED'; row.generated_at=m.now()
    review_service.audit(db,ctx,'report_generated','report_exports',row.id,new={'report_type':report_type,'related_entity_id':related_id})
    db.commit(); return row
