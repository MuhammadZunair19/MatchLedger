import io
from types import SimpleNamespace

import pymupdf
import pytest
from app import models as m
from app import seed_milestone4
from app.matching_schemas import MatchRunInput
from app.security import Context, require_permission
from app.services import report_service, review_service
from app.services.matching.matching_service import execute
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from test_matching_service import add_tenant


def context_for(db, organization_id, user_id, role='owner'):
    membership = SimpleNamespace(organization_id=organization_id, role=role)
    return Context(db.get(m.User, user_id), membership, SimpleNamespace())


def test_role_permissions_are_separated():
    finance = SimpleNamespace(membership=SimpleNamespace(role='finance_approver'))
    require_permission(finance, 'approve')
    with pytest.raises(HTTPException) as exc:
        require_permission(finance, 'review')
    assert exc.value.status_code == 403
    viewer = SimpleNamespace(membership=SimpleNamespace(role='viewer'))
    require_permission(viewer, 'view')
    require_permission(viewer, 'export')


def test_exception_resolution_and_approval_blocking(database):
    base, order_id, receipt_id, invoice_id = add_tenant(database, 'review')
    with database() as db:
        user = db.scalar(select(m.User).where(m.User.id == base.user.id))
        ctx = context_for(db, base.org, user.id)
        run_id = execute(db, ctx, MatchRunInput(invoice_id=invoice_id, purchase_order_id=order_id,
            matching_mode='THREE_WAY', receipt_ids=[receipt_id]))
        invoice = db.get(m.SupplierInvoice, invoice_id)
        approval = review_service.approval(db, ctx, invoice)
        with pytest.raises(HTTPException, match='blocking'):
            review_service.decision(db, ctx, invoice, approval, 'PENDING_APPROVAL', None)
        findings = db.scalars(select(m.MatchFinding).where(m.MatchFinding.match_run_id == run_id)).all()
        for finding in findings:
            review_service.update_exception(db, ctx, finding.id, 1, 'RESOLVED', 'exception_resolved',
                'Verified supplier credit and corrected receiving evidence', resolution_type='CORRECTED_OR_EXPLAINED')
        approval = db.scalar(select(m.InvoiceApproval).where(m.InvoiceApproval.invoice_id == invoice_id))
        result = review_service.decision(db, ctx, invoice, approval, 'PENDING_APPROVAL', 'Procurement review complete')
        assert result.status == 'PENDING_APPROVAL'
        stale_version = result.version - 1
        assert db.scalar(select(m.AuditEvent).where(m.AuditEvent.organization_id == ctx.org,
            m.AuditEvent.action == 'invoice_pending_approval')) is not None
        assert stale_version != result.version


def test_stale_matching_result_is_rejected(database):
    base, order_id, receipt_id, invoice_id = add_tenant(database, 'stale')
    with database() as db:
        ctx = context_for(db, base.org, base.user.id)
        execute(db, ctx, MatchRunInput(invoice_id=invoice_id, purchase_order_id=order_id,
            matching_mode='THREE_WAY', receipt_ids=[receipt_id]))
        invoice = db.get(m.SupplierInvoice, invoice_id)
        approval = review_service.approval(db, ctx, invoice)
        invoice.updated_at = m.now()
        db.flush()
        with pytest.raises(HTTPException, match='stale'):
            review_service.decision(db, ctx, invoice, approval, 'PENDING_REVIEW', None)


def sample_snapshot():
    return {'match_run': {'id':'run-001','status':'EXCEPTION','matching_mode':'THREE_WAY','input_fingerprint':'abc'},
        'supplier': {'name':'شركة الرياض للتوريدات'}, 'invoice': {'number':'INV-2026-001','date':'2026-10-01','due_date':'2026-10-31','currency':'SAR'},
        'purchase_order': {'number':'PO-2026-001'}, 'receipts':[{'number':'GR-2026-001'}],
        'summary': {'currency':'SAR','invoice_value':'5500.00','agreed_order_value':'5000.00','supported_value':'4500.00',
            'price_variance':'500.00','quantity_variance':'500.00','difference_against_supported':'1000.00',
            'comparisons':[{'product':'لفات كابلات كهربائية','sku':'CAB-50','po_quantity':'100','received_quantity':'90',
                'invoice_quantity':'100','po_unit_price':'50.00','invoice_unit_price':'55.00','supported_quantity':'90','line_difference':'1000.00'}]},
        'findings':[{'code':'PRICE_MISMATCH','severity':'BLOCKING','expected_value':'50.00','actual_value':'55.00','difference':'5.00',
            'resolution_status':'OPEN','explanation':'Invoice price exceeds the purchase order price.'}],
        'approval': {'status':'PENDING_REVIEW','decided_at':None,'decision_reason':None}, 'decisions':[],
        'document_versions':{}, 'source_versions':{'invoice_updated_at':'2026-10-01T00:00:00+00:00'}}


def test_generated_pdf_and_excel_are_parseable_and_preserve_values():
    snapshot = sample_snapshot()
    pdf = report_service.build_pdf(snapshot, 'منظمة الاختبار', '2026-10-10T12:00:00+00:00')
    reader = pymupdf.open(stream=pdf, filetype='pdf')
    text = '\n'.join(page.get_text() for page in reader)
    assert reader.page_count >= 2
    assert 'INV-2026-001' in text and '1000.00' in text and 'PENDING_REVIEW' in text
    excel = report_service.build_matching_excel(snapshot)
    workbook = load_workbook(io.BytesIO(excel), data_only=False)
    assert workbook.sheetnames == ['Summary','Line Item Comparison','Exceptions','Approval History','Audit References']
    assert workbook['Summary']['B1'].value == 'INV-2026-001'
    assert workbook['Summary']['B11'].value == 1000
    assert workbook['Line Item Comparison']['F2'].data_type == 'n'


def test_spreadsheet_formula_injection_is_escaped():
    assert report_service.safe_text('=HYPERLINK("bad")').startswith("'=")
    assert report_service.safe_text('+SUM(A1:A2)').startswith("'+")
    assert report_service.safe_text('safe') == 'safe'


def test_milestone4_demo_seed_is_repeatable(database, monkeypatch):
    with database() as db:
        user=m.User(email='demo-seed@example.com',name='Demo Owner',password_hash='test')
        organization=m.Organization(name='Demo Seed Organization')
        db.add_all([user,organization]);db.flush()
        db.add(m.Membership(user_id=user.id,organization_id=organization.id,role='owner'));db.commit()
        organization_id=organization.id
    monkeypatch.setattr(seed_milestone4,'SessionLocal',database)
    seed_milestone4.seed(organization_id)
    seed_milestone4.seed(organization_id)
    with database() as db:
        invoices=db.scalars(select(m.SupplierInvoice).where(m.SupplierInvoice.organization_id==organization_id,
            m.SupplierInvoice.number.like('DEMO%'))).all()
        assert len(invoices)==9
        assert len(db.scalars(select(m.MatchRun).where(m.MatchRun.organization_id==organization_id)).all())==9
        rejected=db.scalar(select(m.InvoiceApproval).join(m.SupplierInvoice,
            m.SupplierInvoice.id==m.InvoiceApproval.invoice_id).where(m.SupplierInvoice.number=='DEMO-INV-E-REJECTED'))
        approved=db.scalar(select(m.InvoiceApproval).join(m.SupplierInvoice,
            m.SupplierInvoice.id==m.InvoiceApproval.invoice_id).where(m.SupplierInvoice.number=='DEMO-INV-F-OVERRIDE'))
        assert rejected.status=='REJECTED'
        assert approved.status=='APPROVED'
